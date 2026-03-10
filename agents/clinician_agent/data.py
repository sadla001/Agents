"""
data.py — Data loading layer for the ONCO-2024-INT Clinician Agent.

Two loaders, both initialised once at startup:

  ProtocolRule      — dataclass for one rule row from the xlsx
  load_rules()      — reads Clinician_Protocol_Rules sheet → dict[domain, list]

  PatientRecord     — holds all three data sections for one patient
  PatientDataLoader — reads Clinician – Patient Care sheet → dict[pid, record]
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

import openpyxl

from .config import RULES_SHEET, DATASET_SHEET


# ══════════════════════════════════════════════════════════════════════════════
# PROTOCOL RULES LOADER
# ══════════════════════════════════════════════════════════════════════════════

@dataclass
class ProtocolRule:
    """One row from the Clinician_Protocol_Rules sheet."""
    rule_id:   str
    domain:    str
    priority:  int
    parameter: str
    formula:   str
    operator:  str
    threshold: str
    unit:      str
    condition: str
    action:    str
    severity:  str
    guideline: str

    # Derived fields — extracted from action text at load time
    main_panel_output: str = field(default="", init=False)
    alert_output:      str = field(default="", init=False)
    action_codes:      list[str] = field(default_factory=list, init=False)

    def __post_init__(self):
        self.main_panel_output = _extract_quoted(self.action, prefix="MAIN PANEL")
        self.alert_output      = _extract_quoted(self.action, prefix="ALERT OUTPUT")
        self.action_codes      = _extract_action_codes(self.action)


def _extract_quoted(text: str, prefix: str = "") -> str:
    """Pull the first quoted string after an optional prefix label."""
    if prefix:
        m = re.search(
            rf'{re.escape(prefix)}:\s*["\u201c]([^"\u201d]+)["\u201d]', text
        )
        if m:
            return m.group(1).strip()
    m = re.search(r'["\u201c]([^"\u201d]+)["\u201d]', text)
    return m.group(1).strip() if m else text.split("\n")[0].strip()


def _extract_action_codes(text: str) -> list[str]:
    m = re.search(r'ACTION:\s*(.+)', text, re.IGNORECASE)
    if m:
        return [c.strip() for c in m.group(1).split("+") if c.strip()]
    return []


VALID_DOMAINS = {"SYMPTOM_BURDEN", "TUMOR_MARKERS", "RENAL", "CLINICAL_ALERT"}


def load_rules(path: Path) -> dict[str, list[ProtocolRule]]:
    """
    Load the 24 clinician protocol rules from the xlsx.
    Returns a dict keyed by domain, each value sorted by priority ascending.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[RULES_SHEET]
    rules: dict[str, list[ProtocolRule]] = {d: [] for d in VALID_DOMAINS}

    for row in ws.iter_rows(min_row=4, values_only=True):
        if not row[0] or not str(row[0]).startswith("RULE-"):
            continue
        cells = [(str(v).strip() if v is not None else "") for v in (list(row) + [None]*14)[:14]]
        domain = cells[1].upper()
        if domain not in VALID_DOMAINS:
            continue
        try:
            priority = int(float(cells[2])) if cells[2] else 99
        except ValueError:
            priority = 99

        rules[domain].append(ProtocolRule(
            rule_id=cells[0], domain=domain, priority=priority,
            parameter=cells[3], formula=cells[4], operator=cells[5],
            threshold=cells[6], unit=cells[7], condition=cells[8],
            action=cells[9], severity=cells[10], guideline=cells[12],
        ))

    for d in rules:
        rules[d].sort(key=lambda r: r.priority)
    return rules


# ══════════════════════════════════════════════════════════════════════════════
# PATIENT DATA LOADER
# ══════════════════════════════════════════════════════════════════════════════

class PatientRecord:
    """Holds Section A (observations), B (medications), C (symptoms) for one patient."""

    def __init__(self, patient_id: str):
        self.patient_id   = patient_id
        self.observations: list[dict] = []
        self.medications:  list[dict] = []
        self.symptoms:     list[dict] = []

    def latest_obs(self, loinc: str, encounter: str = "FOLLOWUP") -> Optional[dict]:
        """Return the most recent FOLLOWUP observation for a given LOINC code."""
        rows = [
            r for r in self.observations
            if str(r.get("LOINC_CODE", "")).strip() == loinc
            and str(r.get("ENCOUNTER_TYPE", "")).strip() == encounter
        ]
        return max(rows, key=lambda r: r.get("VISIT_NUMBER", 0)) if rows else None

    def latest_symptom(self) -> Optional[dict]:
        return max(self.symptoms, key=lambda r: r.get("VISIT_NUMBER", 0)) if self.symptoms else None

    def current_med(self) -> Optional[dict]:
        return self.medications[0] if self.medications else None


class PatientDataLoader:
    """
    Reads all three sections of the Clinician – Patient Care sheet
    and stores them as PatientRecord objects keyed by patient ID.
    """

    def __init__(self, path: Path):
        self._records: dict[str, PatientRecord] = {}
        self._load(path)

    def get(self, patient_id: str) -> Optional[PatientRecord]:
        return self._records.get(patient_id)

    def all_ids(self) -> list[str]:
        return sorted(self._records.keys())

    def count(self) -> int:
        return len(self._records)

    def _load(self, path: Path):
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb[DATASET_SHEET]

        section: Optional[str] = None
        headers: list[str] = []
        buckets: dict[str, dict] = {}

        for row in ws.iter_rows(min_row=1, max_row=400, values_only=True):
            if not any(v is not None for v in row):
                continue
            c0 = str(row[0]).strip() if row[0] else ""

            if   "SECTION A" in c0: section = "A"; headers = []
            elif "SECTION B" in c0: section = "B"; headers = []
            elif "SECTION C" in c0: section = "C"; headers = []
            elif c0 == "PATIENT_ID" and section:
                headers = [str(v).strip() if v else "" for v in row]
            elif c0.startswith("PT-") and headers and section:
                row_dict = {headers[i]: row[i] for i in range(min(len(headers), len(row))) if headers[i]}
                pid = c0
                if pid not in buckets:
                    buckets[pid] = {"A": [], "B": [], "C": []}
                buckets[pid][section].append(row_dict)

        for pid, secs in buckets.items():
            r = PatientRecord(pid)
            r.observations = secs["A"]
            r.medications  = secs["B"]
            r.symptoms     = secs["C"]
            self._records[pid] = r
